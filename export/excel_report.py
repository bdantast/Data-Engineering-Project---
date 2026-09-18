import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows


class ExcelReport:
    def __init__(self):
        self.header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
        self.header_fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
        self.header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def generate(self, output_path, dataframes=None, charts_data=None, ai_analysis=""):
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            if dataframes:
                for name, df in dataframes.items():
                    sheet_name = name[:31].replace("/", "-").replace("\\", "-")
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    self._style_sheet(writer, sheet_name, df)
            if charts_data:
                self._add_charts_sheet(writer, charts_data)
            if ai_analysis:
                self._add_analysis_sheet(writer, ai_analysis)
            self._add_summary_sheet(writer, dataframes or {})
        return output_path

    def _style_sheet(self, writer, sheet_name, df):
        ws = writer.sheets[sheet_name]
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_align
            cell.border = self.border
        for row in ws.iter_rows(min_row=2, max_row=min(len(df) + 1, 1000)):
            for cell in row:
                cell.border = self.border
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 3, 30)

    def _add_charts_sheet(self, writer, charts_data):
        ws = writer.create_sheet("Graficos")
        ws.append(["Grafico", "Tipo", "Dados"])
        ws.append([])
        row = 3
        for chart_info in charts_data:
            ws.cell(row=row, column=1, value=chart_info.get("title", ""))
            ws.cell(row=row, column=2, value=chart_info.get("type", ""))
            df = chart_info.get("data")
            if df is not None and not df.empty:
                for r in dataframe_to_rows(df, index=False, header=True):
                    row += 1
                    for c_idx, value in enumerate(r, 1):
                        ws.cell(row=row, column=c_idx + 2, value=value)
                if chart_info.get("type") == "bar":
                    chart = BarChart()
                    chart.title = chart_info["title"]
                    chart.y_axis.title = "Valor"
                    data_ref = Reference(ws, min_col=4, min_row=row - len(df),
                                         max_row=row, max_col=4)
                    cats_ref = Reference(ws, min_col=3, min_row=row - len(df),
                                         max_row=row)
                    chart.add_data(data_ref, titles_from_data=True)
                    chart.set_categories(cats_ref)
                    chart.width = 20
                    chart.height = 12
                    ws.add_chart(chart, f"A{row + 2}")
            row += 3

    def _add_analysis_sheet(self, writer, analysis_text):
        ws = writer.create_sheet("Analise IA")
        ws.column_dimensions["A"].width = 100
        ws.append(["Analise gerada por Inteligencia Artificial"])
        ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="1A237E")
        ws.append([])
        for line in analysis_text.split("\n"):
            ws.append([line])
        ws.append([])
        ws.append([f"Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}"])

    def _add_summary_sheet(self, writer, dataframes):
        ws = writer.create_sheet("Resumo", 0)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 20
        ws.cell(row=1, column=1, value="DataPulse - Resumo do Relatorio")
        ws.cell(row=1, column=1).font = Font(bold=True, size=16, color="1A237E")
        ws.cell(row=2, column=1, value=f"Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}")
        ws.append([])
        ws.append(["Tabela", "Total de Registros"])
        for name, df in dataframes.items():
            ws.append([name, len(df)])
        if len(dataframes) > 1:
            total = sum(len(df) for df in dataframes.values())
            ws.append(["TOTAL", total])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)


excel_report = ExcelReport()
